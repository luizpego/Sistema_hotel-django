from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from io import BytesIO
from zipfile import ZipFile, BadZipFile

from django.core.exceptions import ValidationError
from django.db import transaction
from django.db.models import Sum
from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from core.models import audit
from core.permissions import ADMIN, RECEPTION, allowed
from financeiro.models import Payment
from hospedes.models import Guest
from quartos.models import Room
from quartos.services import save_room
from reservas.models import Reservation

MAX_ROWS = 1000
EXPORT_LIMIT = 20000
HEADERS = {
    "guests": ["nome", "cpf", "telefone", "email", "cidade", "estado", "nascimento", "observacoes"],
    "rooms": ["numero", "categoria", "capacidade", "diaria", "descricao", "ativo"],
}
LABELS = {"guests": "Hóspedes", "rooms": "Quartos", "reservations": "Reservas", "stays": "Hospedagens",
          "payments": "Pagamentos", "financial": "Financeiro", "occupancy": "Ocupação prevista"}


def safe_cell(value):
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@", "\t", "\r")):
        return "'" + value
    if isinstance(value, datetime) and timezone.is_aware(value):
        return timezone.localtime(value).replace(tzinfo=None)
    return value


def sheet(workbook, title, headers, rows, period="", money_columns=()):
    ws = workbook.create_sheet(title[:31])
    ws.append(["Hotel São João Evangelista · " + title])
    ws.append([period or "Cadastro completo · emitido em " + timezone.localtime().strftime("%d/%m/%Y %H:%M")])
    ws.append(headers)
    for row in rows:
        ws.append([safe_cell(value) for value in row])
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:{get_column_letter(len(headers))}{max(ws.max_row, 3)}"
    ws.sheet_view.showGridLines = False
    for cell in ws[3]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="143F3A")
        cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[3].height = 26
    for row in ws.iter_rows(min_row=4):
        for cell in row:
            if isinstance(cell.value, datetime):
                cell.number_format = "dd/mm/yyyy hh:mm"
            elif isinstance(cell.value, date):
                cell.number_format = "dd/mm/yyyy"
            elif cell.column in money_columns:
                cell.number_format = '"R$" #,##0.00'
    for index, header in enumerate(headers, 1):
        width = max([len(str(header)), *(len(str(ws.cell(row, index).value or "")) for row in range(4, min(ws.max_row, 100) + 1))])
        ws.column_dimensions[get_column_letter(index)].width = min(max(width + 3, 15), 42)
    return ws


def response(workbook, filename):
    stream = BytesIO()
    workbook.save(stream)
    result = HttpResponse(stream.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    result["Content-Disposition"] = f'attachment; filename="{filename}.xlsx"'
    result["Cache-Control"] = "no-store"
    return result


def import_template(kind):
    if kind not in HEADERS:
        raise ValidationError("Tipo de planilha inválido.")
    wb = Workbook()
    ws = wb.active
    ws.title = "Dados"
    ws.append(HEADERS[kind])
    ws.freeze_panes = "A2"
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="143F3A")
        ws.column_dimensions[cell.column_letter].width = 24
    instructions = wb.create_sheet("Instruções")
    lines = ["Preencha a aba Dados e mantenha os cabeçalhos exatamente como estão.",
             "Limite: 1.000 registros. Duplicidades são erros; nenhum cadastro existente será sobrescrito.",
             "Use texto para CPF, telefone e número do quarto para preservar zeros à esquerda.",
             "Não utilize fórmulas, macros, links externos ou células começando com =, +, - ou @."]
    lines += (["Nome e telefone obrigatórios; CPF opcional. Nascimento: AAAA-MM-DD. Estado: sigla de duas letras."] if kind == "guests" else [
        "Categorias: single, double, twin, family, suite. Capacidade: inteiro positivo.",
        "Diária: número positivo com até 2 casas decimais. Ativo: sim ou nao. Novos quartos são liberados para uso."])
    for line in lines:
        instructions.append([line])
    instructions.column_dimensions["A"].width = 115
    return response(wb, f"modelo-{kind}")


def build_record(kind, row):
    values = dict(zip(HEADERS[kind], row))
    if kind == "guests":
        birth = values["nascimento"]
        try:
            birth = date.fromisoformat(birth) if birth else None
        except ValueError:
            raise ValidationError("Nascimento: utilize AAAA-MM-DD.")
        return Guest(name=values["nome"], cpf=values["cpf"] or None, phone=values["telefone"],
                     email=values["email"], city=values["cidade"], state=values["estado"].upper(),
                     birth_date=birth, notes=values["observacoes"])
    try:
        capacity = int(values["capacidade"])
        rate = Decimal(values["diaria"].replace(",", "."))
        if not rate.is_finite():
            raise ValueError
    except (ValueError, InvalidOperation):
        raise ValidationError("Capacidade deve ser inteira e a diária deve ser um valor numérico válido.")
    active = values["ativo"].strip().casefold()
    if active not in ("sim", "nao", "não", "true", "false", "1", "0"):
        raise ValidationError("Ativo: utilize sim ou nao.")
    return Room(number=values["numero"], category=values["categoria"], capacity=capacity,
                daily_rate=rate, description=values["descricao"], active=active in ("sim", "true", "1"))


def validate_rows(kind, rows):
    errors, seen = [], set()
    for number, row in rows:
        try:
            record = build_record(kind, row)
            record.full_clean()
            identities = ([f"cpf:{record.cpf}"] if record.cpf else []) + [f"identity:{record.identity_key}"] if kind == "guests" else [f"number:{record.number}"]
            if any(key in seen for key in identities):
                raise ValidationError("Cadastro duplicado dentro da própria planilha.")
            seen.update(identities)
        except (ValidationError, ValueError, TypeError) as error:
            errors.append({"line": number, "message": " ".join(error.messages) if isinstance(error, ValidationError) else "Dados inválidos. Confira o modelo."})
    return errors


def preview_import(upload, kind):
    if kind not in HEADERS:
        raise ValidationError("Tipo de importação inválido.")
    try:
        raw = upload.read(2 * 1024 * 1024 + 1)
        if len(raw) > 2 * 1024 * 1024:
            raise ValidationError("O arquivo excede 2 MB.")
        with ZipFile(BytesIO(raw)) as archive:
            entries = archive.infolist()
            if len(entries) > 150 or sum(x.file_size for x in entries) > 20 * 1024 * 1024:
                raise ValidationError("A planilha descompactada excede o limite permitido.")
            if any(x.file_size > 10 * 1024 * 1024 or (x.file_size / max(x.compress_size, 1)) > 150 for x in entries):
                raise ValidationError("A compactação da planilha excede os limites de segurança.")
            if any(any(word in x.filename.lower() for word in ("vbaproject", "externallinks", "embeddings", "activex")) for x in entries):
                raise ValidationError("Não são permitidos macros, objetos incorporados ou vínculos externos.")
        wb = load_workbook(BytesIO(raw), read_only=True, data_only=False, keep_links=False)
        try:
            ws = wb["Dados"] if "Dados" in wb.sheetnames else wb.worksheets[0]
            # O limite explícito independe das dimensões declaradas no XML do arquivo.
            ws.reset_dimensions()
            iterator = ws.iter_rows()
            headers = [cell.value for cell in next(iterator)]
            if headers[:len(HEADERS[kind])] != HEADERS[kind] or any(x is not None for x in headers[len(HEADERS[kind]):]):
                raise ValidationError("As colunas não correspondem ao modelo. Baixe o modelo e mantenha os cabeçalhos.")
            rows = []
            for line, cells in enumerate(iterator, 2):
                if line > MAX_ROWS + 1:
                    raise ValidationError("A planilha deve terminar até a linha 1.001, incluindo o cabeçalho.")
                if not any(cell.value is not None for cell in cells):
                    continue
                if any(cell.value is not None for cell in cells[len(HEADERS[kind]):]):
                    raise ValidationError(f"Linha {line}: há colunas extras. Utilize somente as colunas do modelo.")
                if len(rows) >= MAX_ROWS or line > MAX_ROWS + 1:
                    raise ValidationError("A planilha deve conter no máximo 1.000 linhas de dados.")
                row = []
                for cell in cells[:len(HEADERS[kind])]:
                    value = cell.value
                    if cell.data_type == "f" or (isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@"))):
                        raise ValidationError(f"Linha {line}: fórmulas ou conteúdo de fórmula não são permitidos.")
                    if isinstance(value, datetime):
                        value = value.date().isoformat()
                    elif isinstance(value, date):
                        value = value.isoformat()
                    text = str(value) if value is not None else ""
                    if len(text) > 2000:
                        raise ValidationError(f"Linha {line}: célula muito longa (máximo 2.000 caracteres).")
                    row.append(text.strip())
                row.extend([""] * (len(HEADERS[kind]) - len(row)))
                rows.append([line, row])
        finally:
            wb.close()
    except ValidationError:
        raise
    except (BadZipFile, KeyError, OSError, ValueError, StopIteration, IndexError):
        raise ValidationError("Não foi possível ler o arquivo. Envie uma planilha .xlsx válida, sem senha.")
    except Exception as error:
        # Inclui XML inválido ou rejeitado pelo defusedxml, sem revelar detalhes internos.
        raise ValidationError("O arquivo está danificado ou contém conteúdo não permitido.") from error
    if not rows:
        raise ValidationError("A planilha não possui registros para importar.")
    return rows, validate_rows(kind, rows)


@transaction.atomic
def commit_import(kind, rows, actor):
    from django.core.exceptions import PermissionDenied
    if not allowed(actor, ADMIN) and not (kind == "guests" and allowed(actor, RECEPTION)):
        raise PermissionDenied
    errors = validate_rows(kind, rows)
    if errors:
        raise ValidationError([f"Linha {x['line']}: {x['message']}" for x in errors])
    for number, row in rows:
        record = build_record(kind, row)
        record.full_clean()
        record.save()
        audit(actor, "excel.importado", record, line=number)
    return len(rows)


def export_data(kind, start, end, actor):
    from django.core.exceptions import PermissionDenied
    if kind not in LABELS:
        raise ValidationError("Relatório inválido.")
    if not allowed(actor, ADMIN) and not (kind in ("rooms", "reservations", "stays", "guests", "occupancy") and allowed(actor, RECEPTION)):
        raise PermissionDenied
    if (end - start).days > 365 or end < start:
        raise ValidationError("Escolha um período de até 366 dias.")
    wb = Workbook()
    wb.remove(wb.active)
    period = f"Período: {start:%d/%m/%Y} a {end:%d/%m/%Y}"
    def limited(query):
        if query.count() > EXPORT_LIMIT:
            raise ValidationError("O relatório excede 20.000 registros. Reduza o período ou solicite exportação ao responsável técnico.")
        return query.iterator(chunk_size=500)
    if kind == "guests":
        sheet(wb, "Hóspedes", ["Nome", "CPF mascarado", "Telefone", "E-mail", "Cidade", "UF", "Cadastro"],
              ([g.name, g.masked_cpf, g.phone, g.email, g.city, g.state, g.created_at] for g in limited(Guest.objects.all())), money_columns=())
    elif kind == "rooms":
        sheet(wb, "Quartos", ["Número", "Categoria", "Capacidade", "Diária padrão", "Situação operacional", "Ativo"],
              ([r.number, r.get_category_display(), r.capacity, r.daily_rate, r.get_status_display(), "Sim" if r.active else "Não"] for r in limited(Room.objects.all())), money_columns=(4,))
    elif kind in ("reservations", "stays"):
        items = Reservation.objects.filter(start__lte=end, end__gt=start).select_related("guest", "room").annotate(received=Sum("payments__amount"))
        if kind == "stays":
            items = items.filter(status__in=["in_house", "finished"])
        sheet(wb, LABELS[kind], ["Código", "Hóspede", "Quarto", "Entrada", "Saída", "Pessoas", "Diária", "Desconto", "Total", "Recebido", "Pendente", "Situação"],
              ([r.code, r.guest.name, r.room.number, r.start, r.end, r.people, r.daily_rate, r.discount, r.total, r.received or 0, r.total - (r.received or 0), r.get_status_display()] for r in limited(items)), period, (7, 8, 9, 10, 11))
    elif kind in ("payments", "financial"):
        items = Payment.objects.filter(received_at__date__gte=start, received_at__date__lte=end).select_related("reservation", "received_by")
        rows = list(limited(items))
        sheet(wb, "Pagamentos", ["Reserva", "Recebido em", "Valor", "Forma", "Funcionário"],
              ([p.reservation.code, p.received_at, p.amount, p.get_method_display(), str(p.received_by)] for p in rows), period, (3,))
        if kind == "financial":
            daily, monthly, methods = {}, {}, {}
            for p in rows:
                day = timezone.localdate(p.received_at)
                daily[day] = daily.get(day, Decimal(0)) + p.amount
                month = day.strftime("%m/%Y")
                monthly[month] = monthly.get(month, Decimal(0)) + p.amount
                label = p.get_method_display()
                methods[label] = methods.get(label, Decimal(0)) + p.amount
            for title, header, data in [("Receita diária", "Dia", daily), ("Receita mensal", "Mês", monthly), ("Por forma", "Forma", methods)]:
                sheet(wb, title, [header, "Recebido"], sorted(data.items()), period, (2,))
            pending = Reservation.objects.filter(status__in=["pending", "confirmed", "in_house"]).annotate(received=Sum("payments__amount"))
            sheet(wb, "Saldos pendentes", ["Reserva", "Total", "Recebido", "Pendente"],
                  ([r.code, r.total, r.received or 0, r.total - (r.received or 0)] for r in limited(pending) if r.total > (r.received or 0)), "Saldos atuais de todas as reservas ativas; independem do período de recebimentos.", (2, 3, 4))
    elif kind == "occupancy":
        rooms = Room.objects.filter(active=True).count()
        items = list(limited(Reservation.objects.filter(start__lte=end, end__gt=start, status__in=["pending", "confirmed", "in_house", "finished"])))
        rows = []
        day = start
        while day <= end:
            booked = len({r.room_id for r in items if r.start <= day < r.end})
            rows.append([day, rooms, booked, Decimal(booked) / rooms if rooms else Decimal(0)])
            day += timedelta(days=1)
        ws = sheet(wb, "Ocupação prevista", ["Dia", "Quartos ativos hoje", "Quartos reservados", "Taxa prevista"], rows,
                   period + " · conforme datas contratadas; inventário atual de quartos.")
        for row in range(4, ws.max_row + 1):
            ws.cell(row, 4).number_format = "0.0%"
    audit(actor, "excel.exportado", actor, report=kind, start=start.isoformat(), end=end.isoformat())
    return response(wb, f"hotel-{kind}-{start}-{end}")
