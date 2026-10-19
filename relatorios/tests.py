from datetime import timedelta
from decimal import Decimal
from io import BytesIO
from unittest.mock import patch
from django.core.exceptions import ValidationError, PermissionDenied
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import Workbook, load_workbook
from core.testutils import HotelTestCase
from hospedes.models import Guest
from quartos.models import Room
from .excel import HEADERS, LABELS, commit_import, export_data, import_template, preview_import


class ExcelTests(HotelTestCase):
    def upload(self, rows, kind="guests", headers=None):
        wb = Workbook()
        ws = wb.active
        ws.title = "Dados"
        ws.append(headers or HEADERS[kind])
        for row in rows:
            ws.append(row)
        stream = BytesIO()
        wb.save(stream)
        return SimpleUploadedFile("dados.xlsx", stream.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    def test_guest_preview_then_commit(self):
        upload = self.upload([["Joana Excel", "", "33999994444", "", "Belo Horizonte", "MG", "1990-04-01", ""]])
        rows, errors = preview_import(upload, "guests")
        self.assertEqual(errors, [])
        self.assertFalse(Guest.objects.filter(name="Joana Excel").exists())
        self.assertEqual(commit_import("guests", rows, self.reception), 1)
        self.assertTrue(Guest.objects.filter(name="Joana Excel").exists())

    def test_room_import(self):
        rows, errors = preview_import(self.upload([["103", "double", "2", "230.50", "Com varanda", "sim"]], "rooms"), "rooms")
        self.assertEqual(errors, [])
        commit_import("rooms", rows, self.admin)
        self.assertEqual(Room.objects.get(number="103").daily_rate, Decimal("230.50"))

    def test_reception_cannot_import_rooms_even_using_service(self):
        with self.assertRaises(PermissionDenied):
            commit_import("rooms", [], self.reception)

    def test_duplicate_cpf_reports_line(self):
        rows, errors = preview_import(self.upload([["Duplicada", self.guest.cpf, "33999994444"]]), "guests")
        self.assertEqual(errors[0]["line"], 2)
        with self.assertRaises(ValidationError):
            commit_import("guests", rows, self.admin)

    def test_duplicate_in_same_file_rejected(self):
        row = ["Mesmo cadastro", "", "33999994444"]
        rows, errors = preview_import(self.upload([row, row]), "guests")
        self.assertEqual(errors[0]["line"], 3)

    def test_atomic_import_does_not_partially_save(self):
        rows, errors = preview_import(self.upload([["Válida", "", "33999994444"], ["Inválida", "11111111111", "1"]]), "guests")
        self.assertTrue(errors)
        with self.assertRaises(ValidationError):
            commit_import("guests", rows, self.admin)
        self.assertFalse(Guest.objects.filter(name="Válida").exists())

    def test_formula_and_formula_text_rejected(self):
        for value in ['=HYPERLINK("https://example.com")', "+SUM(A1)", "@malicioso"]:
            with self.subTest(value=value), self.assertRaises(ValidationError):
                preview_import(self.upload([[value, "", "33999994444"]]), "guests")

    def test_bad_headers_and_non_xlsx_rejected(self):
        with self.assertRaises(ValidationError):
            preview_import(self.upload([["Dado"]], headers=["errado"]), "guests")
        with self.assertRaises(ValidationError):
            preview_import(SimpleUploadedFile("falso.xlsx", b"not a zip"), "guests")

    def test_row_limit_and_extra_columns(self):
        with self.assertRaises(ValidationError):
            preview_import(self.upload([[str(i), "", "33999994444"] for i in range(1001)]), "guests")
        with self.assertRaises(ValidationError):
            preview_import(self.upload([["Joana", "", "33999994444", "", "", "", "", "", "", "EXTRA"]]), "guests")

    def test_export_all_kinds_are_valid_workbooks(self):
        self.book()
        for kind in LABELS:
            with self.subTest(kind=kind):
                result = export_data(kind, self.today, self.today + timedelta(days=3), self.admin)
                wb = load_workbook(BytesIO(result.content))
                self.assertGreaterEqual(len(wb.sheetnames), 1)
                self.assertTrue(wb.worksheets[0].auto_filter.ref)
                self.assertEqual(wb.worksheets[0].freeze_panes, "A4")

    def test_export_masks_cpf_and_neutralizes_formula(self):
        self.guest.name = "=HYPERLINK(1)"
        self.guest.save()
        result = export_data("guests", self.today, self.today, self.admin)
        ws = load_workbook(BytesIO(result.content)).worksheets[0]
        self.assertEqual(ws["A4"].data_type, "s")
        self.assertTrue(ws["A4"].value.startswith("'="))
        self.assertNotEqual(ws["B4"].value, self.guest.cpf)

    def test_template_has_expected_headers(self):
        result = import_template("rooms")
        wb = load_workbook(BytesIO(result.content))
        self.assertEqual([c.value for c in wb["Dados"][1]], HEADERS["rooms"])

    def test_import_http_preview_and_confirmation(self):
        self.client.force_login(self.reception)
        upload = self.upload([["Importada pela tela", "", "33999995555"]])
        response = self.client.post("/relatorios/importar/", {"kind": "guests", "file": upload})
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Confirmar importação")
        self.assertFalse(Guest.objects.filter(name="Importada pela tela").exists())
        token = self.client.session["excel_preview"]["token"]
        self.assertEqual(self.client.post("/relatorios/importar/confirmar/", {"token": token}).status_code, 302)
        self.assertTrue(Guest.objects.filter(name="Importada pela tela").exists())
        self.assertNotIn("excel_preview", self.client.session)

    def test_expired_or_invalid_confirmation_cannot_import(self):
        self.client.force_login(self.admin)
        self.client.post("/relatorios/importar/confirmar/", {"token": "inventado"})
        self.assertEqual(Guest.objects.count(), 1)

    def test_export_permissions(self):
        with self.assertRaises(PermissionDenied):
            export_data("financial", self.today, self.today, self.reception)
        with self.assertRaises(PermissionDenied):
            export_data("guests", self.today, self.today, self.cleaner)
